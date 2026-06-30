from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_prd import INK, MUTED, add_table, bullet, heading, para, set_font


OUT = Path("施工企业全流程运营托管数字化系统_MVP范围冻结说明书.docx")


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.15

    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(hp.add_run("MVP 范围冻结说明书"), size=9, color=MUTED)

    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(fp.add_run("施工企业全流程运营托管数字化系统"), size=9, color=MUTED)


def title(doc: Document) -> None:
    para(doc, "", after=8)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("MVP 范围冻结说明书"), size=23, color=RGBColor(0, 0, 0), bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    set_font(p.add_run("施工企业全流程运营托管数字化系统"), size=15, color=RGBColor(55, 55, 55))

    for label, value in [
        ("文档目的:", "冻结第一期开发范围，作为原型、数据模型、API、研发排期和验收依据"),
        ("范围建议:", "Web SaaS MVP，优先打通企业/项目底座、资料、成本、财税、回款、预警、驾驶舱闭环"),
        ("评审对象:", "业务负责人、产品经理、研发负责人、测试负责人、实施/托管团队"),
        ("文档版本:", "V1.0"),
        ("生成日期:", date.today().isoformat()),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + " ")
        set_font(r, size=11, color=INK, bold=True)
        set_font(p.add_run(value), size=11, color=INK)

    rule = doc.add_paragraph()
    p_pr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    pbdr.append(bottom)
    p_pr.append(pbdr)


def build() -> None:
    doc = Document()
    setup(doc)
    title(doc)

    heading(doc, "1. MVP 定位", 1)
    para(
        doc,
        "MVP 首期目标不是一次性开发九大模块的完整能力，而是先验证施工企业运营托管系统的核心闭环：企业与项目数据统一、资料证据可沉淀、成本财税可核算、回款清欠可跟进、风险预警可闭环、老板驾驶舱可看见经营结果。",
    )
    para(
        doc,
        "首期建议采用 Web SaaS 管理后台形态，面向企业内部管理人员和托管服务团队使用。移动端、深度 AI、复杂第三方平台对接、军工涉密全量合规能力可后置，但首期必须预留权限、密级、审计和文件安全底座。",
    )

    heading(doc, "2. MVP 目标与成功标准", 1)
    add_table(
        doc,
        ["目标", "成功标准"],
        [
            ("经营闭环", "从项目建档、资料上传、预算支出、财税记录、债权回款到驾驶舱统计可以端到端跑通。"),
            ("风险闭环", "资质到期、资料逾期、支出超预算、无票支出、债权逾期至少五类预警可以自动生成、分配、处理、关闭。"),
            ("证据闭环", "项目合同、资料、票据、债权证据、催收记录、审批记录均可归档并按项目检索。"),
            ("权限闭环", "老板、总经理、项目负责人、财务、资料员、清欠专员、管理员之间有基础权限区分。"),
            ("报表闭环", "项目利润、待回款金额、成本偏差、资料完整度、资质临期、风险事项可在驾驶舱展示并下钻。"),
        ],
        [1800, 7560],
    )

    heading(doc, "3. 首期做什么", 1)
    add_table(
        doc,
        ["模块", "MVP 功能范围", "优先级"],
        [
            ("系统底座", "租户企业、组织部门、用户账号、角色权限、菜单权限、操作日志、消息通知基础能力", "P0"),
            ("项目中心", "项目建档、项目列表、项目详情、甲方信息、合同额、工期、负责人、涉密标识、项目状态", "P0"),
            ("文件中心", "文件上传、分类、项目关联、版本、预览/下载、密级标记、操作日志", "P0"),
            ("预警中心", "预警规则配置、自动触发、责任人、截止日期、处理记录、关闭状态、消息提醒", "P0"),
            ("资质合规", "资质台账、证书附件、有效期、年审日期、到期提醒、核验记录、状态管理", "P0"),
            ("工程资料", "资料目录、资料模板、项目资料上传、审核、缺项预警、竣工组卷基础导出", "P0"),
            ("项目成本", "预算科目、预算金额、支出登记、超预算预警、分包付款记录、成本偏差统计", "P0"),
            ("财税账本", "项目收支流水、发票登记、无票支出标记、税费字段、项目利润表基础版", "P0"),
            ("回款清欠", "债权建档、证据附件、催收阶段、沟通日志、催收文书模板、回款记录、逾期分级", "P0"),
            ("驾驶舱", "项目数、合同额、收入、成本、利润、待回款、风险预警、资料完整度、资质临期总览", "P0"),
        ],
        [1600, 6560, 1200],
        font_size=8.8,
    )

    heading(doc, "4. 首期不做什么", 1)
    add_table(
        doc,
        ["暂不做范围", "原因", "后续阶段"],
        [
            ("完整招投标自动抓取", "外部数据源、解析规则、合规采集和匹配算法复杂，首期可先手工录入投标项目", "V1.1"),
            ("完整标书自动生成", "模板变量、行业差异、版式要求复杂，首期保留模板中心基础能力", "V1.1"),
            ("人脸识别考勤", "涉及设备或第三方 SDK，首期可先手工/导入考勤摘要", "V1.2"),
            ("电子签章深度集成", "需确定供应商、合同、费用和回调机制，首期可上传签署文件", "V1.1"),
            ("发票真伪实时校验", "需税务或第三方接口，首期先做发票记录和异常标记", "V1.2"),
            ("军工涉密完整合规体系", "需确认部署形态、密级边界、保密制度和合规顾问意见", "专项版本"),
            ("并购/分立/转让全流程", "交易流程复杂且低频，首期仅保留资质/人员/财税/债权数据沉淀", "V2.0"),
            ("移动端 App", "首期以 Web 管理后台验证流程，移动端后续按高频现场场景设计", "V1.2"),
            ("复杂 BI 自定义报表", "首期先提供固定指标和筛选下钻，避免报表平台化过早", "V2.0"),
        ],
        [2400, 4660, 2300],
        font_size=8.8,
    )

    heading(doc, "5. 核心业务流程", 1)
    add_table(
        doc,
        ["流程", "关键动作", "MVP 输出"],
        [
            ("企业初始化", "创建企业、组织、角色、用户、项目分类、文件分类、预警规则", "可进入系统并具备基础权限和配置"),
            ("项目建档", "创建项目、关联甲方、合同金额、负责人、工期、涉密标识", "项目成为资料、成本、财税、回款主线"),
            ("资料沉淀", "创建资料目录、上传合同/签证/变更/竣工资料、审核资料", "资料完整度、缺项预警、证据归档"),
            ("成本财税记录", "录入预算、支出、付款、发票、无票支出、税费", "项目利润和风险可计算"),
            ("回款清欠", "建立债权、上传证据、记录沟通、生成文书、登记回款", "债权状态、逾期等级、回款进度可追踪"),
            ("经营查看", "驾驶舱汇总项目利润、成本偏差、待回款、资料完整度、资质临期、风险事项", "老板和运营负责人可以下钻追责"),
        ],
        [1600, 4500, 3260],
    )

    heading(doc, "6. 角色范围", 1)
    add_table(
        doc,
        ["角色", "MVP 权限范围"],
        [
            ("系统管理员", "企业配置、组织用户、角色权限、字典配置、预警规则、日志查看。"),
            ("企业老板", "查看全局驾驶舱、项目经营、利润、待回款、重大风险；审批高风险事项。"),
            ("总经理/运营负责人", "查看全部项目，分配责任人，跟进预警，查看报表和项目详情。"),
            ("项目负责人", "维护项目资料、预算执行、支出申请、资料缺项处理、项目风险处理。"),
            ("财务人员", "维护收支流水、发票、税费、项目利润、回款入账、财税风险处理。"),
            ("资料员", "维护资料目录、上传资料、审核流转、资料缺项处理、组卷导出。"),
            ("清欠专员/法务", "维护债权、证据链、催收文书、沟通日志、回款阶段和清欠预警。"),
            ("资质专员", "维护资质证书、年审、延期、核验记录、到期预警处理。"),
        ],
        [2100, 7260],
        font_size=9,
    )

    heading(doc, "7. MVP 页面清单", 1)
    add_table(
        doc,
        ["一级菜单", "页面", "说明"],
        [
            ("工作台", "首页驾驶舱、我的待办、预警列表", "进入系统后的经营和任务入口"),
            ("项目中心", "项目列表、项目详情、新增/编辑项目", "所有业务围绕项目展开"),
            ("资质合规", "资质列表、资质详情、新增/编辑资质、到期预警", "管理企业资质资产"),
            ("工程资料", "资料目录、资料上传、资料审核、缺项清单、组卷导出", "沉淀结算和回款证据"),
            ("成本管理", "预算列表、支出登记、分包付款、成本偏差", "控制项目盈利"),
            ("财税账本", "收支流水、发票台账、税费记录、项目利润表", "形成项目财税口径"),
            ("回款清欠", "债权列表、债权详情、催收日志、文书模板、回款记录", "提升现金流和证据链完整度"),
            ("文件中心", "文件列表、文件详情、上传文件、版本记录", "全模块附件统一管理"),
            ("系统设置", "组织用户、角色权限、字典配置、预警规则、操作日志", "支撑后台管理"),
        ],
        [1700, 3200, 4460],
        font_size=8.8,
    )

    heading(doc, "8. 里程碑建议", 1)
    add_table(
        doc,
        ["阶段", "周期", "产出"],
        [
            ("第 0 阶段：范围确认", "3-5 天", "确认本说明书，冻结 MVP 功能、不做清单和验收口径"),
            ("第 1 阶段：原型与数据设计", "1-2 周", "页面原型、字段清单、ER 图、权限矩阵"),
            ("第 2 阶段：技术设计", "1 周", "架构方案、API 文档、部署方案、接口方案"),
            ("第 3 阶段：MVP 开发", "6-10 周", "完成底座、项目、资质、资料、成本、财税、回款、驾驶舱"),
            ("第 4 阶段：试点验证", "2-4 周", "导入真实项目数据，验证流程、权限、预警和报表"),
        ],
        [2000, 1500, 5860],
    )

    heading(doc, "9. MVP 验收条件", 1)
    for item in [
        "至少可以完整创建一个企业、三个角色、一个项目，并围绕项目录入资料、预算、支出、发票、债权和回款。",
        "资质到期、资料缺项、支出超预算、无票支出、债权逾期五类预警可以触发、处理、关闭。",
        "文件上传、下载、版本、密级、操作日志可用，关键文件可以按项目和业务类型检索。",
        "驾驶舱可以展示项目利润、待回款、成本偏差、资料完整度、资质临期和风险事项，并支持下钻到明细。",
        "不同角色登录后菜单和数据范围不同，至少具备查看、编辑、导出、审批四类权限控制。",
        "核心列表支持搜索、筛选、分页、导出；新增/编辑表单具备必填校验和错误提示。",
    ]:
        bullet(doc, item)

    heading(doc, "10. 待确认问题", 1)
    add_table(
        doc,
        ["问题", "默认假设", "需要你确认"],
        [
            ("首期客户形态", "先服务单企业/多项目的中小施工企业", "是否需要一开始支持集团多分公司？"),
            ("部署方式", "先按普通 SaaS 设计，涉密能力只预留权限和审计", "是否明确要求私有化或本地部署？"),
            ("移动端", "首期不做 App，只做 Web 管理后台", "现场人员是否必须手机上传资料？"),
            ("招投标", "首期不做自动抓取，只保留后续扩展入口", "是否要把投标项目手工台账纳入 MVP？"),
            ("电子签/短信/OCR", "首期可预留接口，不做深度集成", "哪些第三方能力必须首期上线？"),
            ("财税深度", "首期做项目账本和风险标记，不做税务申报", "是否需要对接现有财务软件？"),
        ],
        [2000, 3860, 3500],
        font_size=9,
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
