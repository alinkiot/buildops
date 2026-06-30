from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_prd import INK, MUTED, add_table, bullet, heading, para, set_font


OUT = Path("施工企业全流程运营托管数字化系统_核心页面原型说明书.docx")


page_specs = [
    {
        "menu": "工作台",
        "page": "首页驾驶舱",
        "purpose": "老板和运营负责人进入系统后的全局经营入口。",
        "layout": "顶部指标卡 + 风险待办 + 项目经营排行 + 回款风险 + 资质临期 + 快捷入口。",
        "filters": "企业、项目、时间范围、项目负责人、甲方、风险等级。",
        "columns": "指标卡：在建项目、合同额、收入、成本、利润、待回款、风险数、资质临期数。",
        "actions": "查看明细、导出报表、进入项目、处理预警、生成经营报告。",
        "states": "正常、关注、预警、高风险。",
    },
    {
        "menu": "工作台",
        "page": "我的待办/预警列表",
        "purpose": "集中处理资质到期、资料缺项、超预算、无票支出、债权逾期等任务。",
        "layout": "筛选区 + 待办列表 + 批量处理 + 详情抽屉。",
        "filters": "来源模块、风险等级、责任人、截止日期、处理状态、项目。",
        "columns": "预警编号、来源模块、标题、项目、风险等级、责任人、截止日期、状态、创建时间。",
        "actions": "查看、处理、转派、关闭、升级、批量导出。",
        "states": "待处理、处理中、已关闭、已逾期、已升级。",
    },
    {
        "menu": "项目中心",
        "page": "项目列表",
        "purpose": "统一管理所有在建、完工和清欠项目。",
        "layout": "筛选区 + 项目表格 + 新增按钮 + 批量导入/导出。",
        "filters": "项目状态、项目类型、甲方、负责人、地域、是否涉密、时间范围。",
        "columns": "项目编号、项目名称、甲方、合同额、负责人、开工日期、计划竣工、状态、待回款、风险数。",
        "actions": "新增项目、编辑、查看详情、导入、导出、归档。",
        "states": "筹备、在建、停工、竣工、结算中、清欠中、归档。",
    },
    {
        "menu": "项目中心",
        "page": "项目详情",
        "purpose": "项目级业务主页面，承载资料、成本、财税、回款和风险的下钻。",
        "layout": "项目基础信息 + 指标卡 + 标签页：资料、成本、财税、回款、文件、预警、日志。",
        "filters": "标签页内按时间、类型、状态筛选。",
        "columns": "基础字段：项目名称、甲方、合同额、负责人、工期、状态、涉密标识、备注。",
        "actions": "编辑项目、上传文件、新增支出、新增债权、处理预警、导出项目报告。",
        "states": "同项目列表状态；涉密项目显示密级标识。",
    },
    {
        "menu": "资质合规",
        "page": "资质列表/详情",
        "purpose": "管理企业资质证书、年审、延期、核验和到期预警。",
        "layout": "资质列表 + 到期日历 + 资质详情 + 附件与核验记录。",
        "filters": "资质类型、证书状态、到期区间、责任人、核验状态。",
        "columns": "资质名称、类型、等级、证书编号、有效期、年审日期、责任人、状态、预警。",
        "actions": "新增资质、编辑、上传证书、核验记录、延期、停用、导出。",
        "states": "草稿、待核验、有效、即将到期、已过期、整改中、停用。",
    },
    {
        "menu": "工程资料",
        "page": "资料目录/上传审核",
        "purpose": "按项目和阶段管理施工资料、签证变更、竣工资料和证据文件。",
        "layout": "左侧资料目录树 + 中部文件列表 + 右侧详情/审核抽屉。",
        "filters": "项目、阶段、资料类型、专业、审核状态、责任人、是否缺项。",
        "columns": "资料名称、阶段、类型、版本、上传人、审核人、截止日期、状态、密级。",
        "actions": "上传、批量上传、审核通过、退回、标记缺项、组卷导出、查看日志。",
        "states": "待上传、待审核、退回、已通过、缺项、已组卷、已归档。",
    },
    {
        "menu": "成本管理",
        "page": "预算与支出",
        "purpose": "建立项目预算并控制支出、分包付款和成本偏差。",
        "layout": "项目筛选 + 预算科目表 + 支出流水 + 超预算预警。",
        "filters": "项目、成本科目、支出类型、供应商/分包商、审批状态、时间范围。",
        "columns": "科目、预算金额、已发生、剩余额度、偏差率；支出编号、金额、对象、状态。",
        "actions": "新增预算、调整预算、新增支出、发起审批、导入、导出。",
        "states": "预算草稿、已发布、执行中、超预算、待审批、已付款、已关闭。",
    },
    {
        "menu": "财税账本",
        "page": "收支流水/发票台账",
        "purpose": "按项目记录收入、成本、发票、税费和财税风险。",
        "layout": "项目账本摘要 + 收支流水 + 发票列表 + 风险提示。",
        "filters": "项目、收支类型、是否有票、发票状态、税率、时间范围。",
        "columns": "流水编号、项目、类型、金额、税率、发票号、对方主体、入账状态、风险标签。",
        "actions": "新增流水、上传发票、标记无票、核销、导出利润表。",
        "states": "待入账、已入账、待补票、异常、待核销、已核销、已归档。",
    },
    {
        "menu": "回款清欠",
        "page": "债权列表/债权详情",
        "purpose": "管理进度款、尾款、质保金、垫资款等债权和催收证据链。",
        "layout": "债权列表 + 分级统计 + 债权详情标签页：证据、催收、文书、回款、日志。",
        "filters": "甲方、项目、欠款类型、逾期区间、催收阶段、风险等级、责任人。",
        "columns": "债权编号、项目、甲方、欠款类型、金额、逾期天数、阶段、责任人、状态。",
        "actions": "新增债权、上传证据、生成文书、记录沟通、登记回款、核销、导出。",
        "states": "正常回款、逾期、呆滞、坏账风险、非诉处理中、部分回款、已结清、已核销。",
    },
    {
        "menu": "文件中心",
        "page": "文件列表/文件详情",
        "purpose": "统一承载全系统附件、版本、密级和操作日志。",
        "layout": "文件分类树 + 文件表格 + 预览区 + 版本/日志标签。",
        "filters": "项目、业务类型、文件类型、密级、上传人、时间范围。",
        "columns": "文件名、业务类型、项目、版本、大小、密级、上传人、上传时间、下载次数。",
        "actions": "上传、预览、下载、替换版本、修改密级、删除、查看日志。",
        "states": "有效、已替换、已归档、已删除、受限访问。",
    },
    {
        "menu": "系统设置",
        "page": "组织用户/角色权限/规则配置",
        "purpose": "配置企业组织、账号、角色、菜单按钮权限、字典和预警规则。",
        "layout": "设置导航 + 配置列表 + 新增/编辑弹窗 + 权限树。",
        "filters": "组织、角色、状态、规则类型、启用状态。",
        "columns": "用户/角色/规则名称、所属组织、权限范围、状态、创建人、更新时间。",
        "actions": "新增、编辑、停用、重置密码、配置权限、启用规则、查看日志。",
        "states": "启用、停用、待激活、已锁定。",
    },
]


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(hp.add_run("核心页面原型说明书"), size=9, color=MUTED)

    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(fp.add_run("施工企业全流程运营托管数字化系统"), size=9, color=MUTED)


def title(doc: Document) -> None:
    para(doc, "", after=8)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("核心页面原型说明书"), size=23, color=RGBColor(0, 0, 0), bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    set_font(p.add_run("施工企业全流程运营托管数字化系统 MVP"), size=15, color=RGBColor(55, 55, 55))
    for label, value in [
        ("文档目的:", "定义 MVP 阶段核心页面、页面区块、字段、按钮、状态和跳转关系"),
        ("原型形式:", "文字原型/开发说明，后续可转低保真或高保真视觉稿"),
        ("适用对象:", "产品经理、UI 设计、前端开发、后端开发、测试"),
        ("文档版本:", "V1.0"),
        ("生成日期:", date.today().isoformat()),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + " ")
        set_font(r, size=11, color=INK, bold=True)
        set_font(p.add_run(value), size=11, color=INK)

    rule = doc.add_paragraph()
    p_pr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    pbdr.append(bottom)
    p_pr.append(pbdr)


def build() -> None:
    doc = Document()
    setup(doc)
    title(doc)

    heading(doc, "1. 信息架构", 1)
    add_table(
        doc,
        ["一级菜单", "二级页面"],
        [
            ("工作台", "首页驾驶舱、我的待办、预警列表"),
            ("项目中心", "项目列表、项目详情、新增/编辑项目"),
            ("资质合规", "资质列表、资质详情、新增/编辑资质、到期预警"),
            ("工程资料", "资料目录、资料上传、资料审核、缺项清单、组卷导出"),
            ("成本管理", "预算与支出、分包付款、成本偏差"),
            ("财税账本", "收支流水、发票台账、税费记录、项目利润表"),
            ("回款清欠", "债权列表、债权详情、催收日志、文书模板、回款记录"),
            ("文件中心", "文件列表、文件详情、版本记录、操作日志"),
            ("系统设置", "组织用户、角色权限、字典配置、预警规则、操作日志"),
        ],
        [1800, 7560],
    )

    heading(doc, "2. 全局交互规则", 1)
    for item in [
        "所有列表页默认包含：关键词搜索、常用筛选、分页、列配置、导出、刷新。",
        "所有新增/编辑表单必须包含：必填校验、格式校验、保存草稿、提交、取消、错误提示。",
        "所有详情页采用“基础信息 + 业务标签页 + 操作日志”的结构，便于追溯。",
        "所有风险事项进入预警中心，并在来源业务页面同步显示处理状态。",
        "所有导出、删除、密级变更、权限变更、审批动作必须写入操作日志。",
        "涉密或敏感数据默认不在列表中展示完整内容，需进入详情或经授权查看。",
    ]:
        bullet(doc, item)

    heading(doc, "3. 页面原型明细", 1)
    rows = []
    for spec in page_specs:
        rows.append(
            (
                spec["menu"],
                spec["page"],
                spec["purpose"],
                spec["layout"],
                spec["filters"],
                spec["columns"],
                spec["actions"],
                spec["states"],
            )
        )
    add_table(
        doc,
        ["菜单", "页面", "页面目的", "布局/区块", "筛选项", "字段/列", "操作按钮", "状态"],
        rows,
        [900, 1150, 1500, 1700, 1400, 1900, 1600, 1210],
        font_size=7.1,
    )

    heading(doc, "4. 关键页面操作路径", 1)
    add_table(
        doc,
        ["路径", "步骤", "结果"],
        [
            ("创建项目", "项目中心 -> 新增项目 -> 填写基础信息 -> 保存 -> 进入项目详情", "项目进入业务主线，可关联资料、成本、财税、回款。"),
            ("处理资料缺项", "预警列表 -> 资料缺项 -> 查看详情 -> 上传资料 -> 提交审核 -> 关闭预警", "资料完整度提升，预警关闭并留痕。"),
            ("超预算支出审批", "成本管理 -> 新增支出 -> 系统判断超预算 -> 发起审批 -> 审批通过后入账", "支出受控，成本偏差更新。"),
            ("建立债权并催收", "回款清欠 -> 新增债权 -> 上传证据 -> 生成文书 -> 记录沟通 -> 登记回款", "债权状态和回款进度可追踪。"),
            ("查看经营风险", "首页驾驶舱 -> 点击风险指标 -> 进入预警列表 -> 下钻到项目/债权/资质", "老板可从总览追到责任人和处理记录。"),
        ],
        [2000, 4900, 2460],
        font_size=9,
    )

    heading(doc, "5. 表单设计原则", 1)
    add_table(
        doc,
        ["表单类型", "设计要求"],
        [
            ("项目表单", "字段分组为基础信息、甲方信息、合同工期、负责人、涉密标识、备注附件。"),
            ("资质表单", "证书编号、资质类型、等级、有效期、年审日期、附件、责任人必须清晰。"),
            ("资料上传", "支持项目、阶段、资料类型、文件、版本、密级、备注；上传后进入审核或直接归档。"),
            ("支出登记", "项目、成本科目、金额、对象、付款日期、附件、是否有票、审批状态必须记录。"),
            ("债权表单", "项目、甲方、欠款类型、金额、到期日、证据附件、责任人、催收阶段必须记录。"),
        ],
        [1800, 7560],
    )

    heading(doc, "6. 待确认交互问题", 1)
    add_table(
        doc,
        ["问题", "默认设计", "需确认"],
        [
            ("首页使用人", "默认老板/总经理优先，展示经营指标", "是否需要为项目负责人、财务、资料员配置不同首页？"),
            ("移动端上传", "MVP 先 Web 上传", "是否必须支持手机端上传现场资料？"),
            ("审批方式", "MVP 使用系统内审批流", "是否需要企业微信/飞书/微信通知审批？"),
            ("文件预览", "支持常见 Office/PDF 图片预览，CAD 先作为附件下载", "CAD 是否必须在线预览？"),
            ("报表导出", "MVP 支持 Excel/PDF 导出", "是否有固定企业模板？"),
        ],
        [2100, 3960, 3300],
        font_size=9,
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
