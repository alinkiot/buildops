from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("施工企业全流程运营托管数字化系统_PRD.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 31, 31)
MUTED = RGBColor(96, 108, 120)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALL_OUT = "F4F6F9"
WHITE = "FFFFFF"


modules = [
    {
        "id": "M01",
        "name": "资质合规管理",
        "position": "合规准入与资质资产底座",
        "pain": "资质、证书、年审、军工保密条件分散管理，导致投标废标、年审遗漏、升级申报缺少数据支撑。",
        "features": [
            "资质资产总台账：建筑施工资质、安许证、劳务资质、军工二级保密资质、涉密信息集成资质、特种作业资质统一建档。",
            "证书原件扫描、线上真伪核验、到期倒计时、等级/承接范围/地域限制智能标注。",
            "年审、延期、抽查、公示、信用扣分、黑名单、企业异常信息预警。",
            "资质升级、增项、并购/分立/转让测算，输出业绩、人员、设备、场地缺口和整改清单。",
            "涉密资质专属管控：政审、社保、培训、场地、保密设备、制度、外协单位资质审核。",
        ],
        "objects": "资质证书、企业主体、人员证书、业绩、设备、场地、保密制度、外协单位、预警事件",
        "priority": "P0",
    },
    {
        "id": "M02",
        "name": "财税风控管理",
        "position": "项目账本与税务风险闭环",
        "pain": "工程款、材料款、劳务费混账，无票支出和税负异常难以前置发现，项目盈亏只能事后核算。",
        "features": [
            "一项目一账本，独立核算收入、直接成本、规费、税费、净利润。",
            "垫资、逾期利息、违约补偿、质保金、坏账、欠款核销专项记账。",
            "无票支出、发票真伪异常、税负率超标、异地施工预缴、简易计税、分包差额计税自动校验。",
            "非诉/诉讼回款专项资金台账，回款税费自动测算。",
            "单项目利润表、企业盈亏报表、税负分析、财税合规体检报告自动生成。",
        ],
        "objects": "项目账本、收支流水、发票、税费规则、回款记录、坏账记录、报表",
        "priority": "P0",
    },
    {
        "id": "M03",
        "name": "招投标全过程管理",
        "position": "机会发现、标书生产与投标复盘",
        "pain": "标讯分散、资料匹配易错、保证金无人跟进、废标原因缺少复盘，涉密投标资料缺少权限管控。",
        "features": [
            "全网标讯抓取，覆盖房建、市政、公路、政府采购、军工涉密项目。",
            "按资质等级、经营范围、地域和项目类型筛选可投项目并推送。",
            "民用建筑、市政、机电、军工涉密工程标书模板库，人员/资质/业绩一键填充。",
            "报名、购标、答疑、开标、保证金、履约保证金全流程节点跟踪。",
            "废标原因归档、中标率、竞品投标数据分析，涉密投标资料加密存储和分级授权。",
        ],
        "objects": "标讯、投标项目、标书模板、企业资料、保证金、竞品记录、废标原因",
        "priority": "P1",
    },
    {
        "id": "M04",
        "name": "工程资料全生命周期管理",
        "position": "资料模板、过程留痕与竣工组卷",
        "pain": "施工资料纸质零散，签证/变更/隐蔽资料滞后，竣工组卷工作量大，涉密图纸泄露风险高。",
        "features": [
            "开工报审、施工记录、隐蔽验收、现场签证、设计变更、竣工验收模板库。",
            "按土建、市政、机电、军工涉密工程区分专业资料模板。",
            "施工计划联动资料填报提醒，签证和变更逾期未上传预警。",
            "资料云端加密存储，上传、下载、转发、打印全程日志留痕。",
            "竣工资料一键组卷，支持 PDF、CAD 图纸批量导出归档，涉密文件分级管控。",
        ],
        "objects": "资料模板、项目阶段、工程文件、签证单、变更单、图纸、操作日志、组卷包",
        "priority": "P0",
    },
    {
        "id": "M05",
        "name": "项目成本精准管控",
        "position": "盈利核心与预算执行控制",
        "pain": "事前预算不准、事中支出失控、分包超付、材料浪费、机械闲置、垫资资金成本无法核算。",
        "features": [
            "人工费、材料费、机械费、管理费、规费、税金一键预算，普通/军工项目差异化模板。",
            "每笔支出对标预算，超预算触发预警和二次审批。",
            "材料领用、机械台班、工人考勤自动联动项目成本。",
            "分包合同、月度结算、付款记录闭环台账，识别超进度付款和重复结算。",
            "自动计算垫资时长、资金占用利息、回款滞后成本，核算项目真实净利润。",
        ],
        "objects": "预算、成本科目、支出单、材料领用、机械台班、考勤、分包合同、付款、利润模型",
        "priority": "P0",
    },
    {
        "id": "M06",
        "name": "劳务与法务一体化",
        "position": "现场用工合规、合同风控与证据链",
        "pain": "实名制、考勤、工资、工伤保险和合同条款缺失，导致劳资纠纷、工伤赔付和维权困难。",
        "features": [
            "工人信息建档、人脸识别考勤、班组管理、电子劳动合同签署。",
            "工伤保险到期提醒，工资台账、银行代发记录、工资公示表一键生成。",
            "总包、分包、材料采购、劳务用工标准合同模板。",
            "合同霸王条款、付款违约、工期漏洞智能识别并标注。",
            "工期逾期、甲方付款逾期、违约行为自动记录，提前预警纠纷并留存证据链。",
        ],
        "objects": "工人、班组、考勤、劳动合同、保险、工资单、工程合同、风险条款、证据",
        "priority": "P1",
    },
    {
        "id": "M07",
        "name": "工程回款清欠专项",
        "position": "债权建档、非诉催收与回款经营分析",
        "pain": "进度款、尾款、质保金、垫资款证据分散，线下催收无流程、无记录、无策略，坏账率高。",
        "features": [
            "进度款、竣工尾款、质保金、垫资款、逾期利息、违约金债权统一建档。",
            "关联合同、竣工验收单、签证单、结算书、沟通记录和送达凭证。",
            "按金额、逾期时长、甲方信用分级为正常回款、呆滞账款、坏账。",
            "催告函、律师函标准化生成，商务谈判、对账、沟通日志全程存档。",
            "回款成功率、平均回款周期、坏账率、甲方信用评级和后续合作风险自动分析。",
        ],
        "objects": "债权、欠款主体、合同证据、催收文书、送达凭证、沟通日志、回款计划、信用评级",
        "priority": "P0",
    },
    {
        "id": "M08",
        "name": "企业数字化驾驶舱",
        "position": "老板视角的经营总控中心",
        "pain": "负责人无法统一查看项目、资金、资质、风险、投标、回款、经营报告，决策依赖经验。",
        "features": [
            "展示在建项目、待回款金额、资质到期预警、财税风险、项目中标率。",
            "资金流水、垫资总额、坏账风险、盈利/亏损项目、高风险甲方可视化。",
            "月度、季度、年度经营分析报告自动生成。",
            "按企业老板、总经理、项目负责人、财务、资料员等角色分级权限。",
            "涉密项目数据和企业营收核心数据独立加密隔离。",
        ],
        "objects": "经营指标、项目汇总、资金汇总、风险事件、分析报告、角色权限、数据看板",
        "priority": "P0",
    },
    {
        "id": "M09",
        "name": "人员体系与资质升级并购托管",
        "position": "人员证书资产、升级测算与并购迁移",
        "pain": "建造师、八大员、技工、特种工、涉密人员证书和社保注册流程混乱，升级/并购缺少资产盘点。",
        "features": [
            "建造师、职称人员、技工、特种工、涉密专职保密人员统一建档。",
            "证书注册到期、继续教育、证件有效期、社保合规自动预警。",
            "按资质升级标准自动核算人员专业、数量、业绩缺口。",
            "跟踪证书聘用、人员补充、资料补录进度，输出补人清单。",
            "并购前资质、人员、业绩、财税、欠款风险线上体检，并购后数据迁移和体系重建。",
        ],
        "objects": "人员、证书、社保、继续教育、资质标准、补人任务、并购体检、迁移记录",
        "priority": "P1",
    },
]


roles = [
    ("企业老板/董事长", "查看经营驾驶舱、审批关键事项、查看涉密/营收核心数据、决策并购与托管方案"),
    ("总经理/运营负责人", "统筹项目、成本、投标、回款、风险预警，查看跨部门报表"),
    ("项目负责人", "维护项目计划、成本执行、资料节点、签证变更、分包进度和现场风险"),
    ("财务人员", "维护项目账本、票据、税费、回款、坏账、财税报表和资金台账"),
    ("资料员", "维护资料模板、阶段资料、竣工组卷、图纸归档和文件流转记录"),
    ("资质/人事专员", "维护资质、人员证书、社保、继续教育、年审延期和升级测算"),
    ("投标专员", "筛选标讯、生成标书、管理保证金、复盘废标/中标数据"),
    ("法务/清欠专员", "维护合同风险、催收文书、沟通日志、证据链、非诉/诉讼回款节点"),
    ("涉密管理员", "管理涉密资质、涉密人员、涉密资料权限、导出打印审批和操作审计"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def para(doc, text="", style=None, before=None, after=None, line_spacing=None, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_font(r, size=11, color=INK)
    fmt = p.paragraph_format
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    if line_spacing is not None:
        fmt.line_spacing = line_spacing
    if align is not None:
        p.alignment = align
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_font(r, size=11, color=INK)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    if level == 1:
        set_font(r, size=16, color=BLUE, bold=True)
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_font(r, size=13, color=BLUE, bold=True)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        set_font(r, size=12, color=DARK_BLUE, bold=True)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], LIGHT_GRAY)
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_font(r, size=font_size, color=INK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(val))
            set_font(r, size=font_size, color=INK)
    set_table_width(table, widths)
    para(doc, "", after=4)
    return table


def add_title_block(doc):
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("产品需求文档 PRD")
    set_font(r, size=9, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("施工企业全流程运营托管数字化系统")
    set_font(r, size=9, color=MUTED)

    para(doc, "", after=8)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("产品需求文档")
    set_font(r, size=23, color=RGBColor(0, 0, 0), bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("施工企业全流程运营托管数字化系统")
    set_font(r, size=15, color=RGBColor(55, 55, 55), bold=False)

    rows = [
        ("产品定位:", "建工全链条合规 + 盈利托管一体化 SaaS 系统"),
        ("适用对象:", "建筑总包、专业分包、劳务工程企业、工程服务商"),
        ("业务模式:", "线上数字化 SaaS 系统 + 线下工程全流程托管服务"),
        ("文档版本:", "V1.0"),
        ("生成日期:", date.today().isoformat()),
        ("来源依据:", "《施工企业全流程运营托管数字化系统方案文档》"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(label + " ")
        set_font(r, size=11, color=INK, bold=True)
        r = p.add_run(value)
        set_font(r, size=11, color=INK)
    para(doc, "", after=8)
    rule = doc.add_paragraph()
    p_pr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALL_OUT)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=10.5, color=DARK_BLUE, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)
    set_table_width(table, [9360])
    para(doc, "", after=6)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for s in ["List Bullet", "List Number"]:
        styles[s].font.name = "Calibri"
        styles[s]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[s].font.size = Pt(11)
        styles[s].paragraph_format.space_after = Pt(8)
        styles[s].paragraph_format.line_spacing = 1.167

    add_title_block(doc)

    heading(doc, "1. 产品概述", 1)
    add_callout(
        doc,
        "PRD 摘要",
        "本产品面向中小施工企业，把资质人员、招投标、施工资料、项目成本、财税、劳务法务、工程回款、经营驾驶舱和并购升级托管整合为一套可运营、可预警、可追溯的数字化系统。产品的核心价值不是单点记账或资料管理，而是用同一套数据底座支持合规前置、风险预警、盈利控制和线下托管服务交付。",
    )
    para(
        doc,
        "产品定位为“建工全链条合规 + 盈利托管”一体化 SaaS 系统，服务建筑总包、专业分包、劳务工程企业及工程服务商。系统强调普通建工合规体系与军工保密合规体系兼容，并把工程非诉回款清欠作为差异化核心模块。",
    )
    para(
        doc,
        "商业定位为“线上数字化 SaaS 管理系统 + 线下工程全流程托管实体服务”，用于解决施工企业合规管理难、投标效率低、项目成本失控、劳务纠纷多、工程款长期拖欠、资质升级缺少方案、经营决策缺少数据的问题。",
    )

    heading(doc, "2. 背景与问题定义", 1)
    pain_rows = [
        ("资质合规", "资质台账混乱、证书到期遗忘、保密资质条件不清、升级并购缺少数据支撑", "投标废标、年审风险、资质资产价值无法释放"),
        ("财税与成本", "工程款/材料款/劳务费混账，无票支出、税负异常，项目盈亏事后核算", "税务稽查风险、利润失真、隐性亏损扩大"),
        ("投标与资料", "标讯遗漏、标书重复填报，施工资料、签证、变更、竣工资料滞后", "中标率低、结算依据不足、竣工验收低效"),
        ("劳务法务", "考勤、工资、合同、保险、工伤证据链不完整", "劳资纠纷、工伤赔付、合同维权困难"),
        ("回款清欠", "债权台账和证据资料分散，催收缺少标准流程和持续跟进", "坏账增加、现金流紧张、甲方合作风险无法量化"),
        ("经营管理", "老板无法统一查看项目、资金、风险、投标、回款、人员和资质", "决策依赖经验，无法形成可复制的托管服务能力"),
    ]
    add_table(doc, ["领域", "核心问题", "业务影响"], pain_rows, [1700, 4200, 3460])

    heading(doc, "3. 产品目标与成功指标", 1)
    add_table(
        doc,
        ["目标类型", "目标描述", "建议指标"],
        [
            ("合规目标", "将资质、证书、涉密条件、人员社保、合同和资料纳入前置校验", "关键资质到期漏报率为 0；重大合规预警闭环率 >= 95%"),
            ("经营目标", "实现项目级预算、成本、税费、回款、利润的实时可视化", "项目真实利润可计算覆盖率 >= 90%；超预算支出审批拦截率 >= 95%"),
            ("效率目标", "减少投标、资料组卷、报表生成、催收文书制作的重复人工", "标书基础资料填充耗时降低 50%；竣工组卷耗时降低 40%"),
            ("回款目标", "建立债权证据链、催收流程和甲方信用评级", "逾期债权建档率 >= 95%；回款节点按期跟进率 >= 90%"),
            ("托管目标", "支撑线下托管团队标准化交付和客户续费", "客户月度经营报告按期生成率 >= 95%；托管任务闭环率 >= 90%"),
        ],
        [1500, 4400, 3460],
    )

    heading(doc, "4. 用户角色与权限", 1)
    add_table(doc, ["角色", "核心权限/使用场景"], roles, [2200, 7160])
    para(
        doc,
        "权限原则：默认按企业、项目、模块、数据密级四层授权；涉密项目、军工资料、企业营收核心数据、并购尽调数据需要独立授权、审批和操作审计。",
    )

    heading(doc, "5. 产品范围与分期", 1)
    add_table(
        doc,
        ["阶段", "建设范围", "目标"],
        [
            ("MVP / V1.0", "资质合规、工程资料、项目成本、财税账本、回款清欠、驾驶舱基础版、统一权限与审计", "先打通合规、资料、成本、财税、回款的核心经营闭环"),
            ("V1.1", "招投标、劳务法务、人员证书、资质升级测算、标准合同/标书/资料模板库增强", "扩展投标前端和现场用工法务闭环，降低运营团队手工工作"),
            ("V2.0", "军工涉密全套管控、并购分立托管、智能经营分析、外部平台对接、AI 风险识别", "形成差异化壁垒和高客单价托管服务能力"),
        ],
        [1500, 5200, 2660],
    )
    add_callout(
        doc,
        "范围边界",
        "PRD 默认建设企业级 SaaS 管理系统，不包含线下托管团队的组织架构、销售合同条款、法律诉讼代理服务本身。非诉与诉讼流程在系统中作为任务、证据、文书、资金和状态管理对象，不替代律师专业判断。",
    )

    heading(doc, "6. 总体业务流程", 1)
    flow = [
        ("企业初始化", "录入企业主体、资质、人员、项目、角色权限、历史债权和模板库", "形成企业数据底座"),
        ("投标前准入", "系统校验资质等级、人员证书、业绩、地域、涉密条件", "输出可投项目/不可投原因"),
        ("项目立项", "建立项目账本、预算、资料计划、合同、分包/供应商、回款计划", "项目进入经营闭环"),
        ("施工过程管控", "同步支出、材料、机械、考勤、签证、变更、合同风险和资料上传", "实时预警成本、资料、劳务、法务风险"),
        ("结算与回款", "汇总竣工资料、债权证据、回款税费、催收节点和甲方信用", "形成回款闭环和坏账管理"),
        ("经营复盘", "驾驶舱汇总项目利润、风险、回款、投标、资质、人员和财税数据", "自动生成经营分析与托管建议"),
    ]
    add_table(doc, ["流程阶段", "关键动作", "输出物"], flow, [1900, 4700, 2760])

    heading(doc, "7. 信息架构与模块总览", 1)
    add_table(
        doc,
        ["模块", "定位", "优先级", "核心数据对象"],
        [(m["name"], m["position"], m["priority"], m["objects"]) for m in modules],
        [1900, 2500, 900, 4060],
        font_size=8.8,
    )

    heading(doc, "8. 功能需求明细", 1)
    for idx, m in enumerate(modules, 1):
        heading(doc, f"8.{idx} {m['name']}（{m['id']}）", 2)
        para(doc, f"模块定位：{m['position']}。")
        para(doc, f"问题定义：{m['pain']}")
        for f in m["features"]:
            bullet(doc, f)
        req_rows = [
            (f"{m['id']}-FR-{i:02d}", item.split("：")[0][:22], item, m["priority"], "功能可配置、可追溯；关键动作写入日志；异常状态触发预警或待办")
            for i, item in enumerate(m["features"], 1)
        ]
        add_table(doc, ["需求编号", "需求项", "需求描述", "优先级", "验收要点"], req_rows, [1150, 1500, 3600, 760, 2350], font_size=8.2)

    heading(doc, "9. 数据模型与跨模块关系", 1)
    add_table(
        doc,
        ["核心实体", "关键字段", "关联模块"],
        [
            ("企业主体", "统一社会信用代码、资质范围、经营地域、信用状态、客户等级", "全模块"),
            ("项目", "项目类型、地域、甲方、合同额、工期、预算、成本、回款计划、涉密标识", "投标、资料、成本、财税、劳务、回款、驾驶舱"),
            ("人员", "岗位、证书、社保、考勤、权限、涉密身份、培训记录", "资质、人员体系、劳务、涉密管控"),
            ("合同", "合同类型、主体、金额、付款条款、工期条款、风险标注、证据附件", "成本、法务、回款、财税"),
            ("工程资料", "资料类型、项目阶段、上传人、版本、密级、审批状态、导出记录", "资料、回款、涉密、驾驶舱"),
            ("债权", "欠款类型、金额、逾期天数、证据完整度、催收阶段、回款状态", "回款、财税、驾驶舱"),
            ("预警事件", "来源模块、风险等级、触发规则、责任人、截止时间、闭环状态", "全模块"),
        ],
        [1500, 4700, 3160],
        font_size=9,
    )

    heading(doc, "10. 规则、预警与审批", 1)
    add_table(
        doc,
        ["规则类别", "触发条件示例", "系统动作"],
        [
            ("资质/证书", "到期前 N 天、年审节点、信用扣分、黑名单、涉密条件缺失", "生成预警、责任人待办、投标准入拦截、整改清单"),
            ("成本预算", "单笔支出超预算、累计成本超阈值、分包超进度付款、重复结算", "二次审批、预算差异报告、项目负责人/老板通知"),
            ("财税", "无票支出超额、发票异常、税负率偏离、异地施工预缴缺失", "风险标记、补票待办、合规体检报告"),
            ("资料", "阶段资料未上传、签证/变更逾期、竣工组卷资料缺项", "资料员待办、结算依据风险提示"),
            ("回款", "逾期天数达到阈值、甲方多次违约、证据缺项、催收节点超期", "债权分级、催收任务、文书生成、升级提醒"),
            ("涉密", "涉密资料导出/打印/转发、外协单位资质未审核、涉密人员培训缺失", "审批拦截、审计日志、涉密管理员告警"),
        ],
        [1550, 4550, 3260],
        font_size=9,
    )

    heading(doc, "11. 非功能需求", 1)
    add_table(
        doc,
        ["类别", "需求"],
        [
            ("安全", "支持多租户数据隔离、RBAC 权限、项目级授权、涉密数据加密、操作审计、导出/打印审批。"),
            ("合规", "关键证据、合同、资料、催收文书需保留版本、时间戳、操作人和流转记录。"),
            ("性能", "常用列表 3 秒内返回；驾驶舱核心指标 5 秒内刷新；批量导入需提供异步任务进度。"),
            ("可用性", "支持 Web 端优先；常用录入表单支持模板导入、批量上传、字段校验和错误提示。"),
            ("可扩展", "预警规则、审批流、模板库、字段字典、角色权限应支持企业级配置。"),
            ("可靠性", "核心业务数据每日备份；关键任务失败可重试；文件上传需校验完整性。"),
            ("审计", "登录、查看、下载、导出、审批、删除、权限变更、涉密文件操作全部留痕。"),
        ],
        [1500, 7860],
    )

    heading(doc, "12. 外部集成需求", 1)
    add_table(
        doc,
        ["集成方向", "说明", "优先级"],
        [
            ("工商/信用/黑名单信息", "获取企业异常、公示、信用扣分、黑名单等风险信息。", "P1"),
            ("住建/招标平台", "采集标讯、资质标准、行业公告、抽查信息。", "P1"),
            ("电子签章", "劳动合同、催告函、律师函、授权书等文书签署和送达。", "P1"),
            ("OCR/文件识别", "识别证书、发票、合同、结算书、验收单、签证单关键字段。", "P1"),
            ("财务/税务接口", "发票校验、税负分析、财务凭证/报表对接。", "P2"),
            ("短信/微信/邮件", "到期、审批、催收、预警消息通知。", "P0"),
        ],
        [1900, 5960, 1500],
    )

    heading(doc, "13. 验收标准", 1)
    add_table(
        doc,
        ["验收维度", "验收口径"],
        [
            ("业务闭环", "MVP 至少完成企业初始化、项目立项、资料上传、预算控制、项目账本、债权建档、预警闭环、驾驶舱汇总。"),
            ("数据一致", "同一项目的合同额、预算、支出、税费、回款、利润在财税、成本、回款、驾驶舱中口径一致。"),
            ("权限安全", "不同角色只能访问授权企业、项目、模块和密级；涉密资料导出/打印必须可审批、可追溯。"),
            ("预警有效", "资质到期、成本超支、资料逾期、回款逾期、无票支出等 P0 规则可以自动触发待办。"),
            ("模板可用", "资质、资料、合同、标书、催收文书等模板可维护、可复用、可导出。"),
            ("报表可用", "单项目利润、企业盈亏、税负分析、回款分析、经营驾驶舱指标可按时间和项目筛选。"),
        ],
        [1700, 7660],
    )

    heading(doc, "14. 风险与待确认问题", 1)
    for item in [
        "军工保密资质和涉密资料管控涉及具体法规、客户实际密级制度及部署环境，需要法务/保密顾问确认系统边界。",
        "标讯抓取、官方资质核验、信用黑名单等外部数据源需要确认可用接口、授权方式、更新频率和合规采集方式。",
        "财税规则存在地域、项目类型、计税方式差异，需要与财税专家共同维护规则库。",
        "人脸识别考勤、电子签章、发票校验、OCR 等能力建议通过成熟第三方服务集成，避免 MVP 阶段自研过重。",
        "线下托管服务的 SOP、人员分工、服务等级和交付物需要另行沉淀，否则系统无法完整承接“托管”业务闭环。",
    ]:
        bullet(doc, item)

    doc.add_page_break()
    heading(doc, "15. 附录：差异化卖点", 1)
    add_table(
        doc,
        ["卖点", "产品化表达"],
        [
            ("普通建工 + 军工保密合规双体系", "通过资质、涉密人员、场地设备、制度、资料密级、外协审核和操作审计形成差异化准入能力。"),
            ("工程非诉回款清欠流程内置", "从债权建档、证据链、催告函/律师函、谈判记录、对账、回款进度到甲方信用评级形成闭环。"),
            ("全生命周期经营闭环", "资质人员 -> 投标 -> 资料 -> 成本 -> 财税 -> 劳务法务 -> 回款 -> 驾驶舱贯通。"),
            ("软件 + 线下托管服务", "系统沉淀托管团队工作流程、客户数据和交付报告，提升服务标准化与续费能力。"),
            ("面向中小施工企业轻量落地", "优先建设高频刚需和可量化价值模块，减少大型 ERP 式复杂功能。"),
        ],
        [2400, 6960],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
